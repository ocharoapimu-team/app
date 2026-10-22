#!/usr/bin/env python3
"""
process_inform.py — extract the Asia-Pacific, climate-related slice of two INFORM workbooks
into inform_data.json for the wiki microsite's INFORM risk page.

  INFORM Risk Index Mid 2026 (v0.7.3, 1 Sep 2026)  -> current risk, natural-hazard components
  INFORM Climate Change Risk Index (2022)           -> 2050 / 2080 projections, hazard change

Scope: countries INFORM assigns to the OCHA ROAP and ROP regions, plus Afghanistan and Pakistan
(INFORM files them under ROMENA; the wiki's El Niño priority tiers include both).
Scores are copied exactly as published; modelled people counts keep 4 decimals. Nothing is rescaled or recomputed.

Usage (from the site folder):
  python src/process_inform.py --risk "<wiki>/raw/<...>_inform-risk-index-mid-2026_2026-09-01.xlsx" \
                               --cc   "<wiki>/raw/<...>_inform-climate-change-risk-index_2022-11-22.xlsx"
  (writes src/inform_data.json; then rebuild the wiki with src/build_site.py)
"""
import argparse, json
from pathlib import Path
import openpyxl

EXTRA = {"AFG", "PAK"}
RISK_COLS = {  # header in 'INFORM Risk 2026 (a-z)' -> key
    "INFORM RISK": "risk", "RISK CLASS": "cls", "Rank": "rank", "Lack of Reliability (*)": "lor",
    "HAZARD & EXPOSURE": "he", "Natural": "nat", "Earthquake": "eq", "River Flood": "flood", "Tsunami": "tsu",
    "Tropical Cyclone": "tc", "Coastal flood": "cflood", "Drought": "drought", "Epidemic": "epi", "Human": "human",
    "VULNERABILITY": "vu", "Socio-Economic Vulnerability": "sev", "Vulnerable Groups": "vg", "Food Security": "food",
    "Recent Shocks": "shocks", "LACK OF COPING CAPACITY": "lcc", "Institutional": "inst", "Infrastructure": "infra",
}


def r1(v):
    return round(v, 4) if isinstance(v, float) else v


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--risk", required=True); ap.add_argument("--cc", required=True); ap.add_argument("--out", default=str(Path(__file__).resolve().parent / "inform_data.json"))
    a = ap.parse_args()
    wb = openpyxl.load_workbook(a.risk, read_only=True, data_only=True)
    regions = {}
    for r in wb["Regions"].iter_rows(min_row=3, values_only=True):
        if r[1]:
            regions[r[1]] = {"ocha": r[6], "wb_income": r[5], "subregion": r[9]}
    scope = {iso for iso, d in regions.items() if d["ocha"] in ("ROAP", "ROP")} | EXTRA
    rows = list(wb["INFORM Risk 2026 (a-z)"].iter_rows(values_only=True))
    hdr = rows[1]
    idx = {h: i for i, h in enumerate(hdr) if h}
    risk = {}
    for r in rows[3:]:
        if not r[1] or r[1] not in scope:
            continue
        d = {"name": r[0], "iso": r[1], "ocha": regions[r[1]]["ocha"], "subregion": regions[r[1]]["subregion"],
             "income": regions[r[1]]["wb_income"]}
        for h, k in RISK_COLS.items():
            v = r[idx[h]]
            d[k] = None if v in ("x", "", None) else r1(v)
        risk[r[1]] = d
    # population, latest reference year (Population Data sheet)
    prow = list(wb["Population Data"].iter_rows(values_only=True))
    years = prow[2]
    last = max(i for i, y in enumerate(years) if isinstance(y, int))
    for r in prow[3:]:
        if r and r[1] in risk:
            risk[r[1]]["pop"] = r[last]; risk[r[1]]["pop_year"] = years[last]
    # climate-hazard exposure indicators (Indicator Data sheet), raw values as published
    IND = {"Annual Expected Exposed People to River Floods": "x_flood",
           "Annual Expected Exposed People to Cyclone's Wind SS1": "x_tc1",
           "Annual Expected Exposed People to Coastal Floods": "x_cflood",
           "Total affected by Drought": "x_drought_aff",
           "Frequency of Drought events": "x_drought_freq",
           "Agriculture Drought probability": "x_agdrought"}
    irow = list(wb["Indicator Data"].iter_rows(values_only=True))
    ih = {h: i for i, h in enumerate(irow[1]) if h}
    for r in irow[2:]:
        if r and r[1] in risk:
            for h, k in IND.items():
                v = r[ih[h]]
                risk[r[1]][k] = None if v in ("x", "", None) else r1(v)

    # INFORM Climate Change
    wc = openpyxl.load_workbook(a.cc, read_only=True, data_only=True)
    ws = wc.worksheets[0]
    crow = list(ws.iter_rows(values_only=True))
    name2iso = {d["name"]: iso for iso, d in risk.items()}
    alias = {"Korea DPR": ["Korea, DPR", "Democratic People's Republic of Korea"], "Korea Republic of": ["Korea, Republic of", "Republic of Korea"],
             "Lao PDR": ["Lao People's Democratic Republic", "Laos"], "Viet Nam": ["Vietnam"], "Micronesia": ["Micronesia (Federated States of)", "Micronesia, Fed. States"]}
    lookup = {}
    for n, iso in name2iso.items():
        lookup[n] = iso
        for al in alias.get(n, []):
            lookup[al] = iso
    cc, unmatched = {}, []
    for r in crow[5:]:
        if not r[0]:
            continue
        iso = lookup.get(str(r[0]).strip())
        if not iso:
            continue
        cc[iso] = {"name_cc": r[0], "b2022": r[1],
                   "p2050": r[2], "p2050_chg": r1(r[3]), "p2050_gap": r1(r[4]),
                   "o2050": r[5], "o2050_chg": r1(r[6]), "o2050_gap": r1(r[7]),
                   "p2080": r[8], "p2080_chg": r1(r[9]), "p2080_gap": r1(r[10]),
                   "o2080": r[11], "o2080_chg": r1(r[12]), "o2080_gap": r1(r[13]),
                   "flood_he_p": r[14], "flood_n_p": r1(r[15]), "flood_n_o": r1(r[16]),
                   "cflood_he_p": r[17], "cflood_n_p": r1(r[18]), "cflood_n_o": r1(r[19]),
                   "drought_he_p": r[20], "drought_n_p": r1(r[21]), "drought_n_o": r1(r[22]),
                   "epi_he_p": r[23], "epi_n_p": r1(r[24]), "epi_n_o": r1(r[25]),
                   "conflict_he_p": r[26], "conflict_he_o": r[27], "pop_b_m": r1(r[28]), "pop_2050p_m": r1(r[29])}
    for iso in risk:
        if iso not in cc:
            unmatched.append(risk[iso]["name"])
    out = {"risk_release": "INFORM Risk Index Mid 2026, v0.7.3 (1 September 2026)",
           "cc_release": "INFORM Climate Change Risk Index (baseline 2022; projections ≈2050 and ≈2080)",
           "scope": "OCHA ROAP and ROP regions as assigned in INFORM's Regions table, plus Afghanistan and Pakistan",
           "countries": [dict(risk[i], cc=cc.get(i)) for i in sorted(risk, key=lambda i: (risk[i]["rank"] or 999))],
           "cc_unmatched": unmatched}
    Path(a.out).write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(risk)} countries; CC matched {len(cc)}; unmatched: {unmatched}")


if __name__ == "__main__":
    main()
